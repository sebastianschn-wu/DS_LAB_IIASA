"""Read the codebook and the policy sheet, and turn them into model instructions.

- ``load_codebook``: variable definitions and the general rules from the codebook .docx
- ``load_allowed_values``: allowed category values from the sheet's CodingScheme
- ``load_policies``: one dict per policy row of the Data sheet
- ``build_instructions``: the system prompt sent to the model
- ``build_response_schema``: JSON schema that forces the model to use allowed values
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import openpyxl

from constants import (
    ALLOWED_COLUMN_MAP,
    CLOSED_FIELD_NOTE,
    CLOSED_FIELDS,
    CODEBOOK_DOCUMENT_XML,
    CODEBOOK_PROSE_HEADINGS,
    CODEBOOK_RULES_HEADING,
    CODEBOOK_VARIABLE_HEADER,
    COL_SHEET_ROW,
    FIELD_YEARS_AMENDED,
    INSTRUCTION_RULES,
    JSON_KEYS_LABEL,
    MODEL_FIELDS,
    MULTI_VALUE_FIELDS,
    MULTI_VALUE_NOTE,
    OPEN_UNIT_FIELDS,
    OPEN_UNIT_NOTE,
    SHEET_CODING_SCHEME,
    SHEET_DATA,
    VALUE_SEPARATOR,
    VARIABLES_HEADING,
    WORDPROCESSING_NS,
    YEARS_AMENDED_NOTE,
)
from text_utils import as_text


def clean_variable_name(raw: str) -> str:
    """'Policy_scope (Policy scope)' -> 'Policy_scope'."""
    name = raw.split("(", 1)[0].strip()
    return name.split()[0] if name else ""


def space_glued_sentences(text: str) -> str:
    """Word paragraphs joined without spaces ('end.Next') get a space back."""
    return re.sub(r"([.!;:])([A-Za-z])", r"\1 \2", text)


def load_codebook(path: Path) -> tuple[list[str], dict[str, str]]:
    """Return (general rules, {variable name: definition}) from the codebook .docx."""
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read(CODEBOOK_DOCUMENT_XML))
    body = root.find(f"{WORDPROCESSING_NS}body")
    if body is None:
        raise SystemExit(f"Codebook has no document body: {path}")

    prose: list[str] = []
    pending_heading: str | None = None
    definitions: dict[str, str] = {}

    for child in list(body):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            line = "".join(node.text or "" for node in child.iter(f"{WORDPROCESSING_NS}t")).strip()
            if not line:
                continue
            if line in CODEBOOK_PROSE_HEADINGS:
                pending_heading = line
                continue
            if pending_heading:
                prose.append(f"{pending_heading}: {space_glued_sentences(line)}")
                pending_heading = None
            continue
        if tag != "tbl":
            continue
        rows: list[list[str]] = []
        for row in child.findall(f"{WORDPROCESSING_NS}tr"):
            cells = [
                space_glued_sentences("".join(n.text or "" for n in cell.iter(f"{WORDPROCESSING_NS}t")).strip())
                for cell in row.findall(f"{WORDPROCESSING_NS}tc")
            ]
            rows.append(cells)
        if not rows or as_text(rows[0][0] if rows[0] else "") != CODEBOOK_VARIABLE_HEADER:
            continue
        for cells in rows[1:]:
            if len(cells) < 2:
                continue
            name = clean_variable_name(cells[0])
            if name:
                definitions[name] = cells[1]

    if not definitions:
        raise SystemExit(f"Could not read codebook variable definitions from {path}")
    return prose, definitions


def load_allowed_values(workbook_path: Path) -> dict[str, list[str]]:
    """Allowed values per model field, read from the CodingScheme sheet."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = workbook[SHEET_CODING_SCHEME].iter_rows(values_only=True)
        header = [as_text(cell) for cell in next(rows)]
        collected: dict[str, list[str]] = {name: [] for name in header if name}
        for row in rows:
            for name, value in zip(header, row):
                text = as_text(value)
                if name and text:
                    collected[name].append(text)
    finally:
        workbook.close()
    return {
        field: collected[column]
        for column, field in ALLOWED_COLUMN_MAP.items()
        if collected.get(column)
    }


def load_policies(path: Path) -> list[dict[str, str]]:
    """All non-empty rows of the Data sheet; ``sheet_row`` is the Excel row number."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[SHEET_DATA].iter_rows(values_only=True)
        header = [as_text(cell) for cell in next(rows)]
        policies: list[dict[str, str]] = []
        for excel_row, row in enumerate(rows, start=2):
            record = {name: as_text(value) for name, value in zip(header, row) if name}
            if not any(record.values()):
                continue
            record[COL_SHEET_ROW] = str(excel_row)
            policies.append(record)
    finally:
        workbook.close()
    return policies


def build_instructions(prose: list[str], definitions: dict[str, str], allowed: dict[str, list[str]]) -> str:
    """System prompt: coding rules, codebook rules, every variable with its definition and allowed values."""
    lines = [*INSTRUCTION_RULES, "", CODEBOOK_RULES_HEADING]
    lines.extend(f"- {item}" for item in prose)
    lines += ["", VARIABLES_HEADING]
    for field in MODEL_FIELDS:
        definition = definitions.get(field, "")
        entry = f"- {field}: {definition}" if definition else f"- {field}"
        values = allowed.get(field)
        if values:
            rendered = VALUE_SEPARATOR.join(values)
            if field in CLOSED_FIELDS:
                entry += CLOSED_FIELD_NOTE.format(values=rendered)
            elif field in OPEN_UNIT_FIELDS:
                entry += OPEN_UNIT_NOTE.format(values=rendered)
        if field in MULTI_VALUE_FIELDS and field != FIELD_YEARS_AMENDED:
            entry += MULTI_VALUE_NOTE
        if field == FIELD_YEARS_AMENDED:
            entry += YEARS_AMENDED_NOTE
        lines.append(entry)
    lines += ["", JSON_KEYS_LABEL + " " + ", ".join(MODEL_FIELDS)]
    return "\n".join(lines)


def build_response_schema(allowed: dict[str, list[str]]) -> dict[str, Any]:
    """JSON schema for Ollama's ``format`` parameter.

    Closed fields may only contain codebook values (or stay empty); multi-value closed
    fields are arrays of codebook values. Without a schema the model can return any
    label, which the old version only flagged afterwards.
    """
    properties: dict[str, Any] = {}
    for field in MODEL_FIELDS:
        values = allowed.get(field)
        if field in CLOSED_FIELDS and values:
            if field in MULTI_VALUE_FIELDS:
                properties[field] = {"type": "array", "items": {"type": "string", "enum": values}}
            else:
                properties[field] = {"type": "string", "enum": [*values, ""]}
        else:
            properties[field] = {"type": "string"}
    return {"type": "object", "properties": properties, "required": list(MODEL_FIELDS)}
