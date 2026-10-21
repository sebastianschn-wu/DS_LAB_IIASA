#!/usr/bin/env python3
"""Code CDR policies with a local Ollama Qwen model.

Reads each resource link in the CDR policy sheet, fetches the document, and
asks Qwen to code that policy using the codebook. Previously coded category
values are not sent to the model.

A document over ``--max-doc-tokens`` is not sent. The row's CDR relevance
summary (the ``CDR_link`` cell) is sent instead. If that summary is missing,
the policy is left out of the model pass and recorded with ``source_used``
of ``omitted``. The same summary fallback applies when the link cannot be
fetched or the extracted text is empty.

Results are written to ``results/first_loop_coding.csv`` and updated after
every row, so a run can be stopped and continued. Completed rows are skipped
unless ``--force`` is set. Paths, column names, and prompt text live in
``constants.py``.

Usage:
    python src/scripts/first-loop/first-loop.script.py
    python src/scripts/first-loop/first-loop.script.py --limit 5
    python src/scripts/first-loop/first-loop.script.py --max-doc-tokens 4000
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
import time
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urljoin
from xml.etree import ElementTree as ET

_SCRIPT_DIR = str(Path(__file__).resolve().parent)
if _SCRIPT_DIR not in sys.path:
    sys.path.insert(0, _SCRIPT_DIR)

import openpyxl
import requests
from bs4 import BeautifulSoup
from PyPDF2 import PdfReader

from constants import (
    ALLOWED_COLUMN_MAP,
    CACHEABLE_OUTCOMES,
    CLOSED_FIELD_NOTE,
    CLOSED_FIELDS,
    CODEBOOK_DOCUMENT_XML,
    CODEBOOK_PROSE_HEADINGS,
    CODEBOOK_RULES_HEADING,
    CODEBOOK_VARIABLE_HEADER,
    COL_CODED_AT,
    COL_DOCUMENT_TITLE,
    COL_DOCUMENT_TOKENS,
    COL_FALLBACK_REASON,
    COL_FETCH_ERROR,
    COL_ISO,
    COL_JURISDICTION,
    COL_LINK,
    COL_MODEL,
    COL_MODEL_ERROR,
    COL_OMIT_REASON,
    COL_SHEET_ROW,
    COL_SOURCE_USED,
    COL_TOKENS_LOWER_BOUND,
    COL_VALIDATION_NOTES,
    COL_YEAR_ENACTED,
    COMPLETE_SOURCES,
    CSV_ENCODING,
    DEFAULT_CODEBOOK,
    DEFAULT_HOST,
    DEFAULT_MAX_DOC_TOKENS,
    DEFAULT_MODEL,
    DEFAULT_OUTPUT,
    DEFAULT_SHEET,
    DOWNLOAD_TIMEOUT,
    EMPTY_MODEL_VALUES,
    ENCRYPTED_PDF_ERROR,
    FALLBACK_BY_OUTCOME,
    FALLBACK_EMPTY_DOCUMENT,
    FALLBACK_OVER_TOKEN_THRESHOLD,
    FETCH_ATTEMPTS,
    FETCH_ERROR,
    FETCH_OUTCOME,
    FETCH_TEXT,
    FETCH_TOKENS,
    FETCH_TOKENS_LOWER_BOUND,
    FIELD_YEARS_AMENDED,
    HTML_PARSER,
    HTTP_ACCEPT,
    IDENTITY_COLUMNS,
    INSTRUCTION_RULES,
    JSON_KEYS_LABEL,
    JSON_OBJECT_ERROR,
    LABEL_DOCUMENT_TITLE,
    LABEL_ISO,
    LABEL_JURISDICTION,
    LABEL_URL,
    LABEL_YEAR_ENACTED,
    MAX_DOWNLOAD_BYTES,
    MIN_DOCUMENT_TOKENS,
    MISSING_LINK_ERROR,
    MODEL_CALL_ATTEMPTS,
    MODEL_FIELDS,
    MULTI_VALUE_FIELDS,
    MULTI_VALUE_NOTE,
    NOTE_EXPECTED_SINGLE,
    NOTE_UNIT_OUTSIDE,
    NOTE_VALUE_OUTSIDE,
    OLLAMA_CHAT_TIMEOUT,
    OLLAMA_CONTEXT_CAP,
    OLLAMA_CONTEXT_FLOOR,
    OLLAMA_CONTEXT_RESERVE,
    OLLAMA_CONTEXT_SAFETY_FACTOR,
    OLLAMA_KEEP_ALIVE,
    OLLAMA_MISSING_MESSAGE,
    OLLAMA_NUM_PREDICT,
    OLLAMA_PULL_CHECK_TIMEOUT,
    OLLAMA_SEED,
    OLLAMA_START_TIMEOUT_SECONDS,
    OLLAMA_TAGS_TIMEOUT,
    OMIT_REASON_SUFFIX,
    OPEN_UNIT_FIELDS,
    OPEN_UNIT_NOTE,
    OTHER_LABEL,
    OTHER_PREFIX,
    OUTCOME_EMPTY,
    OUTCOME_ERROR,
    OUTCOME_OK,
    OUTCOME_OVER_THRESHOLD,
    PDF_EXTENSION,
    PDF_MAGIC,
    RUN_COLUMNS,
    SHEET_CODING_SCHEME,
    SHEET_DATA,
    SOURCE_CACHE_DIRNAME,
    SOURCE_DOCUMENT,
    SOURCE_ERROR,
    SOURCE_LABEL_DOCUMENT,
    SOURCE_LABEL_SUMMARY,
    SOURCE_OMITTED,
    SOURCE_SUMMARY,
    SUMMARY_COLUMN,
    TEXT_ENCODING,
    TOKENS_LOWER_BOUND_NO,
    TOKENS_LOWER_BOUND_YES,
    USER_AGENT,
    USER_PROMPT_HEADER,
    USER_PROMPT_RETURN,
    USER_PROMPT_SOURCE_TEXT,
    USER_PROMPT_SOURCE_TYPE,
    VALUE_SEPARATOR,
    VARIABLES_HEADING,
    WORDPROCESSING_NS,
    YEARS_AMENDED_NOTE,
)


class DocumentOverThreshold(Exception):
    """The retrieved file is too large to be a usable document under the token cap."""


def estimate_tokens(text: str) -> int:
    """Approximate a Qwen token count.

    English and similar scripts are counted at about four characters per token.
    CJK, kana, and hangul characters are counted as one token each.
    """
    cjk = 0
    other = 0
    for char in text:
        code = ord(char)
        if (
            0x3400 <= code <= 0x9FFF
            or 0x3040 <= code <= 0x30FF
            or 0xAC00 <= code <= 0xD7AF
            or 0xF900 <= code <= 0xFAFF
        ):
            cjk += 1
        else:
            other += 1
    return cjk + (other + 3) // 4


def as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def fold_label(value: str) -> str:
    folded = value.replace("CO₂", "CO2").replace("₂", "2")
    folded = re.sub(r"\s+", " ", folded).strip().casefold()
    return folded


def clean_variable_name(raw: str) -> str:
    name = raw.split("(", 1)[0].strip()
    return name.split()[0] if name else ""


def space_glued_sentences(text: str) -> str:
    return re.sub(r"([.!;:])([A-Za-z])", r"\1 \2", text)


def load_codebook(path: Path) -> tuple[list[str], dict[str, str]]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read(CODEBOOK_DOCUMENT_XML))
    body = root.find(f"{WORDPROCESSING_NS}body")
    if body is None:
        raise SystemExit(f"Codebook has no document body: {path}")

    prose: list[str] = []
    pending_heading: str | None = None
    definitions: dict[str, str] = {}

    for child in list(body):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            line = "".join(node.text or "" for node in child.iter(f"{WORDPROCESSING_NS}t")).strip()
            if not line:
                continue
            if line in CODEBOOK_PROSE_HEADINGS:
                pending_heading = line
                continue
            if pending_heading:
                prose.append(f"{pending_heading}: {space_glued_sentences(line)}")
                pending_heading = None
            continue
        if tag != "tbl":
            continue
        rows: list[list[str]] = []
        for row in child.findall(f"{WORDPROCESSING_NS}tr"):
            cells: list[str] = []
            for cell in row.findall(f"{WORDPROCESSING_NS}tc"):
                parts = [
                    node.text or ""
                    for node in cell.iter(f"{WORDPROCESSING_NS}t")
                ]
                cells.append(space_glued_sentences("".join(parts).strip()))
            rows.append(cells)
        if not rows or as_text(rows[0][0] if rows[0] else "") != CODEBOOK_VARIABLE_HEADER:
            continue
        for cells in rows[1:]:
            if len(cells) < 2:
                continue
            name = clean_variable_name(cells[0])
            if name:
                definitions[name] = cells[1]

    if not definitions:
        raise SystemExit(f"Could not read codebook variable definitions from {path}")
    return prose, definitions


def load_allowed_values(workbook_path: Path) -> dict[str, list[str]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_CODING_SCHEME]
        rows = sheet.iter_rows(values_only=True)
        header = [as_text(cell) for cell in next(rows)]
        collected: dict[str, list[str]] = {name: [] for name in header if name}
        for row in rows:
            for name, value in zip(header, row):
                text = as_text(value)
                if name and text:
                    collected[name].append(text)
    finally:
        workbook.close()

    allowed: dict[str, list[str]] = {}
    for column, field in ALLOWED_COLUMN_MAP.items():
        values = collected.get(column, [])
        if values:
            allowed[field] = values
    return allowed


def load_policies(path: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_DATA]
        rows = sheet.iter_rows(values_only=True)
        header = [as_text(cell) for cell in next(rows)]
        policies: list[dict[str, str]] = []
        for excel_row, row in enumerate(rows, start=2):
            record = {
                name: as_text(value)
                for name, value in zip(header, row)
                if name
            }
            if not any(record.values()):
                continue
            record[COL_SHEET_ROW] = str(excel_row)
            policies.append(record)
    finally:
        workbook.close()
    return policies


def build_instructions(prose: list[str], definitions: dict[str, str], allowed: dict[str, list[str]]) -> str:
    lines = [*INSTRUCTION_RULES, "", CODEBOOK_RULES_HEADING]
    lines.extend(f"- {item}" for item in prose)
    lines.append("")
    lines.append(VARIABLES_HEADING)
    for field in MODEL_FIELDS:
        definition = definitions.get(field, "")
        entry = f"- {field}: {definition}" if definition else f"- {field}"
        values = allowed.get(field)
        if values:
            rendered = VALUE_SEPARATOR.join(values)
            if field in CLOSED_FIELDS:
                entry += CLOSED_FIELD_NOTE.format(values=rendered)
            elif field in OPEN_UNIT_FIELDS:
                entry += OPEN_UNIT_NOTE.format(values=rendered)
        if field in MULTI_VALUE_FIELDS and field != FIELD_YEARS_AMENDED:
            entry += MULTI_VALUE_NOTE
        if field == FIELD_YEARS_AMENDED:
            entry += YEARS_AMENDED_NOTE
        lines.append(entry)
    lines.append("")
    lines.append(JSON_KEYS_LABEL + " " + ", ".join(MODEL_FIELDS))
    return "\n".join(lines)


def cache_path(cache_dir: Path, url: str, max_doc_tokens: int) -> Path:
    digest = hashlib.sha256(f"{max_doc_tokens}|{url}".encode(TEXT_ENCODING)).hexdigest()
    return cache_dir / f"{digest}.json"


def read_cache(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding=TEXT_ENCODING))
    except (OSError, json.JSONDecodeError):
        return None


def write_cache(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding=TEXT_ENCODING)


def decode_bytes(data: bytes, content_type: str) -> str:
    match = re.search(r"charset=([^;\s]+)", content_type, flags=re.IGNORECASE)
    encoding = match.group(1).strip("\"'") if match else TEXT_ENCODING
    try:
        return data.decode(encoding, errors="replace")
    except LookupError:
        return data.decode(TEXT_ENCODING, errors="replace")


def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, HTML_PARSER)
    for tag in soup(["script", "style", "noscript", "svg", "form"]):
        tag.decompose()
    return normalize_text(soup.get_text("\n"))


def first_pdf_link(html: str, base_url: str) -> str:
    soup = BeautifulSoup(html, HTML_PARSER)
    for anchor in soup.find_all("a", href=True):
        href = anchor["href"].strip()
        if PDF_EXTENSION in href.lower():
            return urljoin(base_url, href)
    return ""


def pdf_to_text(data: bytes, max_doc_tokens: int) -> tuple[str, int, bool]:
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError(ENCRYPTED_PDF_ERROR) from exc
    parts: list[str] = []
    tokens = 0
    for page in reader.pages:
        try:
            extracted = page.extract_text() or ""
        except Exception:
            continue
        cleaned = normalize_text(extracted)
        if not cleaned:
            continue
        parts.append(cleaned)
        tokens = estimate_tokens("\n".join(parts))
        if tokens > max_doc_tokens:
            return "", tokens, True
    text = "\n".join(parts)
    return text, estimate_tokens(text), False


def download(session: requests.Session, url: str) -> tuple[bytes, str, str]:
    response = session.get(url, timeout=DOWNLOAD_TIMEOUT, allow_redirects=True, stream=True)
    response.raise_for_status()
    chunks: list[bytes] = []
    total = 0
    for chunk in response.iter_content(chunk_size=65536):
        if not chunk:
            continue
        total += len(chunk)
        if total > MAX_DOWNLOAD_BYTES:
            response.close()
            raise DocumentOverThreshold(
                f"download exceeded {MAX_DOWNLOAD_BYTES} bytes"
            )
        chunks.append(chunk)
    content_type = response.headers.get("Content-Type", "")
    final_url = response.url
    return b"".join(chunks), content_type, final_url


def is_pdf(url: str, content_type: str, data: bytes) -> bool:
    if data.startswith(PDF_MAGIC):
        return True
    if PDF_EXTENSION.lstrip(".") in content_type.lower():
        return True
    path = url.split("?", 1)[0].lower()
    return path.endswith(PDF_EXTENSION)


def fetch_payload(
    outcome: str,
    tokens: int,
    lower_bound: bool,
    text: str,
    error: str = "",
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        FETCH_OUTCOME: outcome,
        FETCH_TOKENS: tokens,
        FETCH_TOKENS_LOWER_BOUND: lower_bound,
        FETCH_TEXT: text,
    }
    if error:
        payload[FETCH_ERROR] = error
    return payload


def fetch_document(
    session: requests.Session,
    url: str,
    max_doc_tokens: int,
    cache_dir: Path,
) -> dict[str, Any]:
    cached = read_cache(cache_path(cache_dir, url, max_doc_tokens))
    if cached and cached.get(FETCH_OUTCOME) in CACHEABLE_OUTCOMES:
        return cached

    def attempt(target: str, allow_pdf_hop: bool) -> dict[str, Any]:
        data, content_type, final_url = download(session, target)
        if is_pdf(final_url, content_type, data):
            text, tokens, over = pdf_to_text(data, max_doc_tokens)
            if over:
                return fetch_payload(OUTCOME_OVER_THRESHOLD, tokens, True, "")
            if tokens < MIN_DOCUMENT_TOKENS:
                return fetch_payload(OUTCOME_EMPTY, tokens, False, "")
            return fetch_payload(OUTCOME_OK, tokens, False, text)

        html = decode_bytes(data, content_type)
        if allow_pdf_hop:
            pdf_url = first_pdf_link(html, final_url)
            if pdf_url and pdf_url != target:
                preview = html_to_text(html)
                if estimate_tokens(preview) < MIN_DOCUMENT_TOKENS:
                    return attempt(pdf_url, allow_pdf_hop=False)
        text = html_to_text(html)
        tokens = estimate_tokens(text)
        if tokens > max_doc_tokens:
            return fetch_payload(OUTCOME_OVER_THRESHOLD, tokens, False, "")
        if tokens < MIN_DOCUMENT_TOKENS:
            return fetch_payload(OUTCOME_EMPTY, tokens, False, "")
        return fetch_payload(OUTCOME_OK, tokens, False, text)

    last_error = ""
    for try_index in range(FETCH_ATTEMPTS):
        try:
            fetched = attempt(url, allow_pdf_hop=True)
            write_cache(cache_path(cache_dir, url, max_doc_tokens), fetched)
            return fetched
        except DocumentOverThreshold:
            fetched = fetch_payload(OUTCOME_OVER_THRESHOLD, max_doc_tokens + 1, True, "")
            write_cache(cache_path(cache_dir, url, max_doc_tokens), fetched)
            return fetched
        except Exception as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            if try_index == 0:
                time.sleep(1.5)
    return fetch_payload(OUTCOME_ERROR, 0, False, "", last_error)


def server_is_up(host: str) -> bool:
    try:
        response = requests.get(f"{host}/api/tags", timeout=OLLAMA_TAGS_TIMEOUT)
        return response.ok
    except requests.RequestException:
        return False


def ensure_ollama(host: str, model: str) -> None:
    if shutil.which("ollama") is None:
        raise SystemExit(OLLAMA_MISSING_MESSAGE)

    if server_is_up(host):
        print(f"Ollama is already running at {host}")
    else:
        print("Starting Ollama locally...")
        subprocess.Popen(
            ["ollama", "serve"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            start_new_session=True,
        )
        deadline = time.time() + OLLAMA_START_TIMEOUT_SECONDS
        while time.time() < deadline:
            if server_is_up(host):
                print("Ollama is ready.")
                break
            time.sleep(0.5)
        else:
            raise SystemExit(f"Ollama did not become ready at {host}")

    response = requests.get(f"{host}/api/tags", timeout=OLLAMA_PULL_CHECK_TIMEOUT)
    response.raise_for_status()
    installed: list[str] = []
    for item in response.json().get("models", []):
        installed.append(item.get("name", ""))
        installed.append(item.get("model", ""))
    wanted = model if ":" in model else f"{model}:latest"
    base = model.split(":", 1)[0]
    present = wanted in installed or model in installed or any(
        name.split(":", 1)[0] == base for name in installed if name
    )
    if present:
        print(f"Using local model {model}")
        return
    print(f"Pulling {model}...")
    subprocess.run(["ollama", "pull", model], check=True)


def parse_json_object(text: str) -> dict[str, Any]:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        payload = json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start < 0 or end <= start:
            raise
        payload = json.loads(cleaned[start : end + 1])
    if not isinstance(payload, dict):
        raise ValueError(JSON_OBJECT_ERROR)
    return payload


def stringify_model_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return VALUE_SEPARATOR.join(
            stringify_model_value(item) for item in value if stringify_model_value(item)
        )
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    text = as_text(value)
    if text.casefold() in EMPTY_MODEL_VALUES:
        return ""
    return text


def canonicalize(field: str, raw: str, allowed: dict[str, list[str]]) -> tuple[str, list[str]]:
    text = stringify_model_value(raw)
    if not text:
        return "", []
    choices = allowed.get(field)
    pieces = [part.strip() for part in text.split(";")] if field in MULTI_VALUE_FIELDS else [text]
    pieces = [part for part in pieces if part]
    if not choices:
        if field == FIELD_YEARS_AMENDED:
            return VALUE_SEPARATOR.join(pieces), []
        return VALUE_SEPARATOR.join(pieces) if field in MULTI_VALUE_FIELDS else text, []

    lookup = {fold_label(choice): choice for choice in choices}
    notes: list[str] = []
    canonical: list[str] = []
    for piece in pieces:
        match = lookup.get(fold_label(piece))
        if match:
            canonical.append(match)
            continue
        other = lookup.get(OTHER_LABEL)
        if other and fold_label(piece).startswith(OTHER_LABEL):
            canonical.append(piece if piece.casefold().startswith(OTHER_LABEL) else f"{OTHER_PREFIX}{piece}")
            continue
        if field in OPEN_UNIT_FIELDS:
            canonical.append(piece)
            notes.append(NOTE_UNIT_OUTSIDE.format(field=field, piece=piece))
            continue
        canonical.append(piece)
        notes.append(NOTE_VALUE_OUTSIDE.format(field=field, piece=piece))

    if field not in MULTI_VALUE_FIELDS and len(pieces) > 1:
        notes.append(NOTE_EXPECTED_SINGLE.format(field=field))
    joined = VALUE_SEPARATOR.join(dict.fromkeys(canonical))
    return joined, notes


def coerce_coding(payload: dict[str, Any], allowed: dict[str, list[str]]) -> tuple[dict[str, str], str]:
    folded = {str(key).strip().casefold(): value for key, value in payload.items()}
    coded: dict[str, str] = {}
    notes: list[str] = []
    for field in MODEL_FIELDS:
        raw = payload.get(field, folded.get(field.casefold(), ""))
        value, field_notes = canonicalize(field, raw, allowed)
        coded[field] = value
        notes.extend(field_notes)
    return coded, VALUE_SEPARATOR.join(notes)


def call_model(
    host: str,
    model: str,
    instructions: str,
    user_prompt: str,
    num_ctx: int,
    temperature: float,
) -> str:
    response = requests.post(
        f"{host}/api/chat",
        json={
            "model": model,
            "stream": False,
            "format": "json",
            "keep_alive": OLLAMA_KEEP_ALIVE,
            "messages": [
                {"role": "system", "content": instructions},
                {"role": "user", "content": user_prompt},
            ],
            "options": {
                "temperature": temperature,
                "seed": OLLAMA_SEED,
                "num_ctx": num_ctx,
                "num_predict": OLLAMA_NUM_PREDICT,
            },
        },
        timeout=OLLAMA_CHAT_TIMEOUT,
    )
    response.raise_for_status()
    body = response.json()
    content = body.get("message", {}).get("content", "")
    if not content:
        raise RuntimeError(f"empty model response: {json.dumps(body)[:400]}")
    return content


def build_user_prompt(policy: dict[str, str], source_label: str, source_text: str) -> str:
    return "\n".join(
        [
            USER_PROMPT_HEADER,
            f"{LABEL_JURISDICTION}: {policy.get(COL_JURISDICTION, '')}",
            f"{LABEL_ISO}: {policy.get(COL_ISO, '')}",
            f"{LABEL_DOCUMENT_TITLE}: {policy.get(COL_DOCUMENT_TITLE, '')}",
            f"{LABEL_YEAR_ENACTED}: {policy.get(COL_YEAR_ENACTED, '')}",
            f"{LABEL_URL}: {policy.get(COL_LINK, '')}",
            "",
            f"{USER_PROMPT_SOURCE_TYPE} {source_label}",
            USER_PROMPT_SOURCE_TEXT,
            source_text,
            "",
            USER_PROMPT_RETURN,
        ]
    )


def blank_result(policy: dict[str, str], model: str) -> dict[str, str]:
    row = {column: "" for column in IDENTITY_COLUMNS + RUN_COLUMNS + MODEL_FIELDS}
    for column in IDENTITY_COLUMNS:
        if column == COL_SHEET_ROW:
            row[column] = policy.get(COL_SHEET_ROW, "")
        else:
            row[column] = policy.get(column, "")
    row[COL_MODEL] = model
    row[COL_CODED_AT] = datetime.now().isoformat(timespec="seconds")
    return row


def choose_source(
    policy: dict[str, str],
    fetched: dict[str, Any],
    max_doc_tokens: int,
) -> tuple[str, str, str, str]:
    """Return source text, source_used, fallback_reason, omit_reason."""
    summary = normalize_text(policy.get(SUMMARY_COLUMN, ""))
    outcome = fetched.get(FETCH_OUTCOME)
    reason = FALLBACK_BY_OUTCOME.get(outcome, "")

    if outcome == OUTCOME_OK and int(fetched.get(FETCH_TOKENS) or 0) <= max_doc_tokens:
        return as_text(fetched.get(FETCH_TEXT)), SOURCE_DOCUMENT, "", ""

    if not reason:
        reason = (
            FALLBACK_OVER_TOKEN_THRESHOLD
            if int(fetched.get(FETCH_TOKENS) or 0) > max_doc_tokens
            else FALLBACK_EMPTY_DOCUMENT
        )

    if summary and estimate_tokens(summary) <= max_doc_tokens:
        return summary, SOURCE_SUMMARY, reason, ""
    return "", SOURCE_OMITTED, "", f"{reason}{OMIT_REASON_SUFFIX}"


def code_policy(
    policy: dict[str, str],
    fetched: dict[str, Any],
    instructions: str,
    allowed: dict[str, list[str]],
    host: str,
    model: str,
    max_doc_tokens: int,
    num_ctx: int,
    temperature: float,
) -> dict[str, str]:
    row = blank_result(policy, model)
    row[COL_DOCUMENT_TOKENS] = str(fetched.get(FETCH_TOKENS) or "")
    row[COL_TOKENS_LOWER_BOUND] = (
        TOKENS_LOWER_BOUND_YES if fetched.get(FETCH_TOKENS_LOWER_BOUND) else TOKENS_LOWER_BOUND_NO
    )
    row[COL_FETCH_ERROR] = as_text(fetched.get(FETCH_ERROR))

    source_text, source_used, fallback_reason, omit_reason = choose_source(
        policy, fetched, max_doc_tokens
    )
    row[COL_SOURCE_USED] = source_used
    row[COL_FALLBACK_REASON] = fallback_reason
    row[COL_OMIT_REASON] = omit_reason
    if source_used == SOURCE_OMITTED:
        return row

    if source_used == SOURCE_DOCUMENT:
        source_label = SOURCE_LABEL_DOCUMENT
    else:
        source_label = SOURCE_LABEL_SUMMARY.format(fallback_reason=fallback_reason)
    prompt = build_user_prompt(policy, source_label, source_text)
    last_error = ""
    content = ""
    for _ in range(MODEL_CALL_ATTEMPTS):
        try:
            content = call_model(host, model, instructions, prompt, num_ctx, temperature)
            coded, notes = coerce_coding(parse_json_object(content), allowed)
            row.update(coded)
            row[COL_VALIDATION_NOTES] = notes
            row[COL_MODEL_ERROR] = ""
            return row
        except Exception as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            time.sleep(1)
    row[COL_MODEL_ERROR] = last_error
    if content:
        row[COL_MODEL_ERROR] += " | response: " + content[:500].replace("\n", " ")
    return row


def result_columns() -> list[str]:
    return IDENTITY_COLUMNS + RUN_COLUMNS + MODEL_FIELDS


def load_existing(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    with path.open(newline="", encoding=CSV_ENCODING) as handle:
        return {row[COL_SHEET_ROW]: row for row in csv.DictReader(handle) if row.get(COL_SHEET_ROW)}


def is_complete(row: dict[str, str]) -> bool:
    return row.get(COL_SOURCE_USED) in COMPLETE_SOURCES and not row.get(COL_MODEL_ERROR)


def write_results(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", newline="", encoding=CSV_ENCODING) as handle:
        writer = csv.DictWriter(handle, fieldnames=result_columns(), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sheet", type=Path, default=DEFAULT_SHEET)
    parser.add_argument("--codebook", type=Path, default=DEFAULT_CODEBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--model",
        default=DEFAULT_MODEL,
        help=f"Local Ollama model name. Default: {DEFAULT_MODEL}",
    )
    parser.add_argument("--host", default=DEFAULT_HOST)
    parser.add_argument(
        "--max-doc-tokens",
        type=int,
        default=DEFAULT_MAX_DOC_TOKENS,
        help=(
            "Maximum estimated tokens of document text sent to the model. "
            f"Default: {DEFAULT_MAX_DOC_TOKENS}"
        ),
    )
    parser.add_argument("--limit", type=int, default=0, help="Code at most this many pending policies. 0 codes all.")
    parser.add_argument("--temperature", type=float, default=0.0)
    parser.add_argument("--force", action="store_true", help="Recode rows already present in the output CSV.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    if args.max_doc_tokens < MIN_DOCUMENT_TOKENS:
        raise SystemExit(f"--max-doc-tokens must be at least {MIN_DOCUMENT_TOKENS}")

    prose, definitions = load_codebook(args.codebook)
    allowed = load_allowed_values(args.sheet)
    policies = load_policies(args.sheet)
    instructions = build_instructions(prose, definitions, allowed)
    instruction_tokens = estimate_tokens(instructions)
    # Extra room covers token-estimate error plus the JSON answer.
    num_ctx = min(
        OLLAMA_CONTEXT_CAP,
        max(
            OLLAMA_CONTEXT_FLOOR,
            int((instruction_tokens + args.max_doc_tokens) * OLLAMA_CONTEXT_SAFETY_FACTOR)
            + OLLAMA_CONTEXT_RESERVE,
        ),
    )

    print(f"Policies in sheet: {len(policies)}")
    print(f"Codebook instruction tokens (estimate): {instruction_tokens}")
    print(f"Document token threshold: {args.max_doc_tokens}")
    print(f"Ollama context window: {num_ctx}")

    ensure_ollama(args.host.rstrip("/"), args.model)

    cache_dir = args.output.parent / SOURCE_CACHE_DIRNAME
    existing = {} if args.force else load_existing(args.output)
    results = {row_id: row for row_id, row in existing.items()}
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept": HTTP_ACCEPT})

    pending = [
        policy
        for policy in policies
        if args.force or not is_complete(results.get(policy[COL_SHEET_ROW], {}))
    ]
    if args.limit:
        pending = pending[: args.limit]
    print(f"Policies to code this run: {len(pending)}")

    for index, policy in enumerate(pending, start=1):
        title = policy.get(COL_DOCUMENT_TITLE, "")
        print(
            f"[{index}/{len(pending)}] row {policy[COL_SHEET_ROW]} "
            f"{policy.get(COL_JURISDICTION, '')} | {title}"
        )
        link = policy.get(COL_LINK, "")
        try:
            if not link:
                fetched = fetch_payload(OUTCOME_ERROR, 0, False, "", MISSING_LINK_ERROR)
            else:
                fetched = fetch_document(session, link, args.max_doc_tokens, cache_dir)
            row = code_policy(
                policy,
                fetched,
                instructions,
                allowed,
                args.host.rstrip("/"),
                args.model,
                args.max_doc_tokens,
                num_ctx,
                args.temperature,
            )
        except Exception as exc:
            row = blank_result(policy, args.model)
            row[COL_SOURCE_USED] = SOURCE_ERROR
            row[COL_MODEL_ERROR] = f"{type(exc).__name__}: {exc}"
        results[policy[COL_SHEET_ROW]] = row
        ordered = [results[item[COL_SHEET_ROW]] for item in policies if item[COL_SHEET_ROW] in results]
        write_results(args.output, ordered)
        print(
            f"    source={row[COL_SOURCE_USED]} tokens={row[COL_DOCUMENT_TOKENS] or '-'} "
            f"fallback={row[COL_FALLBACK_REASON] or '-'} omit={row[COL_OMIT_REASON] or '-'}"
        )
        if link:
            time.sleep(0.4)

    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
